# -*- coding: utf-8 -*-
"""Relatório filtrável de frequência por canonical_term."""

from __future__ import annotations

import pandas as pd
from openpyxl import load_workbook

import textura_freq as tf


def _frame():
    return pd.DataFrame({
        "canonical_term": ["clot", "clot", "clot", "static", "static", "clot"],
        "matched_form": ["clot", "cloth", "cloth", "static", "static", "clot"],
        "caminho_ficheiro": ["a", "a", "b", "x", "y", "a"],
        "doc_id": ["a", "a", "b", "x", "y", "a"],
        "nuclear": [True, True, True, True, True, False],
    })


def test_agrega_e_ordena():
    f = tf.tabela_frequencia(tf.filtrar(_frame(), nuclear="true"))
    # static: 2 hits / 2 docs ; clot: 3 hits / 2 docs (docs a,b) → ordena por N_hits desc
    linhas = {r.canonical_term: (r.N_hits, r.n_documentos) for r in f.itertuples()}
    assert linhas["clot"] == (3, 2)
    assert linhas["static"] == (2, 2)
    assert list(f["canonical_term"]) == ["clot", "static"]  # 3 > 2


def test_filtro_forma_e_nuclear():
    f = tf.tabela_frequencia(tf.filtrar(
        _frame(), nuclear="true", termos=["clot"], formas=["cloth"]))
    assert f.loc[0, "N_hits"] == 2 and f.loc[0, "n_documentos"] == 2
    g = tf.tabela_frequencia(tf.filtrar(_frame(), nuclear="false", termos=["clot"]))
    assert g.loc[0, "N_hits"] == 1 and g.loc[0, "n_documentos"] == 1


def test_reusa_grafico_xlsx(tmp_path):
    out = tf.gerar(_frame(), tmp_path, etiqueta_consulta="teste", nuclear="true")
    ws = load_workbook(out["xlsx"])["Frequência"]
    assert ws["B5"].value == "N_hits" and ws["C5"].value == "n_documentos"
    assert len(ws._charts) == 1 and ws._charts[0].grouping == "clustered"


def test_carregar_base_sync_e_override(tmp_path):
    """Hits velho + revisto_por_humano em Concordancia → nuclear humano."""
    conc = pd.DataFrame({
        "canonical_term": ["clot"],
        "matched_form": ["clot"],
        "doc_id": ["a"],
        "caminho_ficheiro": ["a"],
        "nuclear": [False],
        "revisto_por_humano": [True],
        "relacao_sintactica": ["atributiva"],
    })
    hits = conc.copy()
    hits["revisto_por_humano"] = [""]
    xlsx = tmp_path / "rev.xlsx"
    with pd.ExcelWriter(xlsx) as xw:
        conc.to_excel(xw, sheet_name="8_Concordancia", index=False)
        hits.to_excel(xw, sheet_name="8_Concordancia_Hits", index=False)
    df = tf.carregar_base(xlsx)
    assert bool(df.loc[0, "nuclear"]) is True


def test_familia_prefere_substantivo():
    assert tf.e_substantivo("homogeneity")
    assert not tf.e_substantivo("homogeneous")
    assert not tf.e_substantivo("statically")
    formas = ["homogeneous"] * 10 + ["homogeneity"] * 5 + ["homogenize"] * 2
    assert tf.familia_do_grupo(formas) == "homogeneity"
    assert tf.familia_do_grupo(["static", "static", "statically"]) == "static"


def test_tabela_familia_agrega_hits(tmp_path):
    df = pd.DataFrame({
        "canonical_term": ["homogen"] * 6,
        "matched_form": [
            "homogeneous", "homogeneous", "homogeneity",
            "homogeneity", "homogene", "homogenize",
        ],
        "doc_id": list("abcdef"),
        "nuclear": [True] * 6,
    })
    f = tf.tabela_frequencia_familia(tf.filtrar(df, nuclear="true"))
    assert f.loc[0, "familia"] == "homogeneity"
    assert int(f.loc[0, "N_hits"]) == 6
    assert int(f.loc[0, "n_documentos"]) == 6
    out = tf.gerar(df, tmp_path, etiqueta_consulta="fam", nuclear="true")
    ws = load_workbook(out["xlsx_familia"])["Frequência"]
    assert ws["A5"].value == "familia"
    assert ws["A6"].value == "homogeneity"
    assert int(ws["B6"].value) == 6
    assert len(ws._charts) == 1


def test_familia_do_lexico(tmp_path):
    lex = tmp_path / "campo.txt"
    lex.write_text("homogen:E / homogeneity = homogen*\n", encoding="utf-8")
    mapa = tf.ler_mapa_familia_lexico(lex)
    assert mapa["homogen"] == "homogeneity"
    df = pd.DataFrame({
        "canonical_term": ["homogen", "homogen"],
        "matched_form": ["homogeneous", "homogeneous"],
        "doc_id": ["a", "b"],
        "nuclear": [True, True],
    })
    f = tf.tabela_frequencia_familia(df, lexico=mapa)
    assert f.loc[0, "familia"] == "homogeneity"


def test_so_leitura_nao_escreve_hits(tmp_path):
    conc = pd.DataFrame({
        "canonical_term": ["clot"],
        "matched_form": ["clot"],
        "doc_id": ["a"],
        "nuclear": [True],
        "relacao_sintactica": ["atributiva"],
    })
    hits = pd.concat([conc, conc], ignore_index=True)
    xlsx = tmp_path / "div.xlsx"
    with pd.ExcelWriter(xlsx) as xw:
        conc.to_excel(xw, sheet_name="8_Concordancia", index=False)
        hits.to_excel(xw, sheet_name="8_Concordancia_Hits", index=False)
    df = tf.carregar_base(xlsx, sincronizar=False)
    assert len(df) == 1
    hits2 = pd.read_excel(xlsx, sheet_name="8_Concordancia_Hits")
    assert len(hits2) == 2


def test_comparar_espelho_divergiu(tmp_path):
    import textura_analise as ta
    conc = pd.DataFrame({"nuclear": [True], "hit_key": ["a"]})
    hits = pd.DataFrame({"nuclear": [True, False], "hit_key": ["a", "b"]})
    xlsx = tmp_path / "e.xlsx"
    with pd.ExcelWriter(xlsx) as xw:
        conc.to_excel(xw, sheet_name="8_Concordancia", index=False)
        hits.to_excel(xw, sheet_name="8_Concordancia_Hits", index=False)
    est = ta.comparar_espelho_hits(xlsx)
    assert est["divergiu"] is True
    assert est["n_conc"] == 1 and est["n_hits"] == 2
