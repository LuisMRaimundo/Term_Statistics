# -*- coding: utf-8 -*-
"""Surgical Phase-2 fixes: polarity guard, overrides, sheets, Sankey caps."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

import textura_analise as ta
import textura_plots as tplot


def _wb_minimo(path: Path, *, nuclear=None, extra_sheet=True,
               revisto=None, polaridade=None):
    n = 4
    nuclear = [True, True, True, False] if nuclear is None else nuclear
    conc = pd.DataFrame({
        "hit_key": [f"h{i}" for i in range(n)],
        "canonical_term": ["static", "static", "uniform", "varied"],
        "matched_form": ["static", "static", "uniform", "varied"],
        "nuclear": nuclear,
        "relacao_sintactica": [
            "atributiva", "atributiva", "predicativa", "incidental"],
        "polaridade": polaridade or [
            "estabilidade", "estabilidade", "estabilidade", ""],
        "eixo": ["ambos", "ambos", "ambos", "ambos"],
        "doc_id": ["DocA", "DocA", "DocB", "DocC"],
        "contexto": ["ctx a", "ctx a", "ctx b", "ctx c"],
        "revisto_por_humano": revisto or ["", "", "", ""],
        "caminho_ficheiro": ["a.pdf", "a.pdf", "b.pdf", "c.pdf"],
    })
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        pd.DataFrame({
            "chave": ["consulta", "janelas_kwic_processadas"],
            "valor": ["textur* NEAR/4", 100],
        }).to_excel(xw, sheet_name="0_Instrucoes", index=False)
        conc.to_excel(xw, sheet_name="8_Concordancia", index=False)
        conc.to_excel(xw, sheet_name="8_Concordancia_Hits", index=False)
        if extra_sheet:
            pd.DataFrame({"nota": ["legend"], "n": [1]}).to_excel(
                xw, sheet_name="Resumo_ligacao", index=False)


class TestDeterminismoPolaridade(unittest.TestCase):
    def test_omit_binomial_when_v_is_1(self):
        df = pd.DataFrame({
            "canonical_term": ["static"] * 3 + ["uniform"] * 2,
            "polaridade": ["estabilidade"] * 5,
        })
        det, v = ta.dv_deterministica_de_canonical(df, "polaridade")
        self.assertTrue(det)
        self.assertGreaterEqual(v, 0.999)
        row = ta._linha_teste_omitido("polaridade", "polaridade", v)
        self.assertEqual(row["familia"], "polaridade")
        self.assertIn("deterministic function of canonical_term", row["teste"])
        self.assertIn("Cramér's V = 1", row["teste"])

    def test_relacao_nao_e_omitida(self):
        df = pd.DataFrame({
            "canonical_term": ["static", "static", "uniform", "uniform"],
            "relacao_sintactica": [
                "atributiva", "predicativa", "atributiva", "predicativa"],
        })
        det, _ = ta.dv_deterministica_de_canonical(df, "relacao_sintactica")
        self.assertFalse(det)


class TestOverrideHumano(unittest.TestCase):
    def test_vazio_nao_altera(self):
        conc = pd.DataFrame({
            "hit_key": ["a", "b"],
            "nuclear": [True, False],
            "revisto_por_humano": ["", None],
        })
        out, logs = ta.aplicar_override_revisto(conc)
        self.assertEqual(list(out["nuclear"]), [True, False])
        self.assertEqual(logs, [])

    def test_false_decrementa(self):
        conc = pd.DataFrame({
            "hit_key": ["a"],
            "nuclear": [True],
            "revisto_por_humano": [False],
        })
        out, logs = ta.aplicar_override_revisto(conc)
        self.assertFalse(bool(out.loc[0, "nuclear"]))
        self.assertEqual(len(logs), 1)
        self.assertIn("hit_key=a", logs[0])
        self.assertIn("nuclear_computado=True", logs[0])
        self.assertIn("nuclear_humano=False", logs[0])


class TestSankeyCaps(unittest.TestCase):
    def test_sem_teto_soma_n(self):
        df = pd.DataFrame({
            "matched_form": ["static", "static", "uniform"],
            "doc_id": ["A", "B", "A"],
        })
        pares = tplot.pares_forma_obra(df, max_pares=None)
        self.assertEqual(sum(w for *_, w in pares), 3)

    def test_termo_rel_inclui_polaridade_vazia(self):
        df = pd.DataFrame({
            "canonical_term": ["smooth", "static"],
            "relacao_sintactica": ["atributiva", "atributiva"],
            "polaridade": ["", "estabilidade"],
        })
        pares = tplot.pares_termo_rel_pol(df, max_pares=None)
        camada1 = sum(w for a, b, w in pares if a in {"smooth", "static"})
        self.assertEqual(camada1, 2)
        pols = {b for a, b, w in pares if a == "atributiva"}
        self.assertIn("não adjudicada", pols)
        self.assertIn("estabilidade", pols)


class TestFolhasPreservadas(unittest.TestCase):
    def test_resumo_ligacao_sobrevive(self):
        with tempfile.TemporaryDirectory() as td:
            src = Path(td) / "in.xlsx"
            dst = Path(td) / "out.xlsx"
            _wb_minimo(src)
            ta.analisar(src, dst, desduplicacao="nenhuma")
            with pd.ExcelFile(dst) as xl:
                self.assertIn("Resumo_ligacao", xl.sheet_names)
                self.assertIn("1_Resumo", xl.sheet_names)
            self.assertTrue((dst.parent / "_g_freq_token.xlsx").is_file())
            wb = load_workbook(dst)
            self.assertTrue(wb["6_Graficos_barras"]._charts)
            self.assertIn("6_Graficos_familias", wb.sheetnames)
            self.assertTrue(wb["6_Graficos_familias"]._charts)
            wb.close()
            lig = pd.read_excel(dst, sheet_name="Resumo_ligacao")
            self.assertEqual(lig.loc[0, "nota"], "legend")
            res = pd.read_excel(dst, sheet_name="1_Resumo")
            val = dict(zip(res["indicador"].astype(str), res["valor"]))
            self.assertEqual(int(val["Linhas brutas (hits)"]), 4)
            self.assertEqual(int(val["Linhas nucleares (N_hits)"]), 3)


class TestNuvemFonte(unittest.TestCase):
    def test_freq_nuvem_igual_12_formas(self):
        formas = pd.DataFrame({
            "canonical_term": ["static", "static", "uniform"],
            "matched_form": ["static", "statically", "uniform"],
            "n": [20, 5, 3],
            "n_documentos": [8, 3, 3],
        })
        freq = ta.freq_nuvem_de_formas(formas)
        self.assertEqual(freq["static"], 20)
        self.assertEqual(freq["statically"], 5)
        self.assertEqual(sum(freq.values()), 28)
        docs = ta.freq_nuvem_de_formas(formas, ponderacao="docs")
        self.assertEqual(docs["static"], 8)
        self.assertEqual(docs["statically"], 3)
        self.assertNotEqual(freq, docs)
        fam = ta.mapa_forma_familia(formas)
        self.assertEqual(fam["statically"], "static")
        tab = ta._df_pesos_grafico(freq, ponderacao="hits", familia=fam)
        self.assertEqual(list(tab["forma"]), ["static", "statically", "uniform"])
        self.assertEqual(list(tab["peso"]), [20, 5, 3])
        self.assertTrue((tab["ponderacao"] == "hits").all())
        self.assertEqual(tab.loc[tab["forma"] == "statically", "canonical_term"].iloc[0],
                         "static")

    def test_cor_familia_partilhada(self):
        import textura_plots as tplot
        fam = {"homogeneous": "homogen", "homogenous": "homogen",
               "static": "static"}
        cores = tplot.cores_por_forma(fam)
        self.assertEqual(cores["homogeneous"], cores["homogenous"])
        self.assertEqual(cores["homogeneous"], tplot.cor_canonical("homogen"))
        self.assertNotEqual(cores["homogeneous"], cores["static"])


class TestDispersao(unittest.TestCase):
    def test_gini_e_share(self):
        self.assertAlmostEqual(ta.gini_index([1, 1, 1, 1]), 0.0, places=6)
        self.assertGreater(ta.gini_index([73, 1, 1, 1]), 0.5)
        df = pd.DataFrame({
            "canonical_term": ["static"] * 5,
            "doc_id": ["A"] * 3 + ["B", "C"],
        })
        disp = ta._dispersao_por_chave(df, ["canonical_term"], "doc_id")
        row = disp.iloc[0]
        self.assertEqual(int(row["n_documentos"]), 3)
        self.assertAlmostEqual(float(row["share_doc_maximo"]), 0.6)
        self.assertEqual(row["doc_dominante"], "A")


class TestPlanoEKappa(unittest.TestCase):
    def test_cohen_kappa_acordo_total(self):
        self.assertEqual(ta.cohen_kappa([1, 1, 0, 0], [1, 1, 0, 0]), 1.0)

    def test_plano_prevalece_sobre_cli(self):
        d, n, c, r, logs = ta.aplicar_plano_a_priori(
            {"desduplicacao": "contexto", "relacao": "atributiva"},
            desduplicacao="nenhuma", nulo_polaridade="banda",
            cooc_unidade="obra", relacoes=None)
        self.assertEqual(d, "contexto")
        self.assertEqual(r, ["atributiva"])
        self.assertTrue(logs)

    def test_analisar_escreve_kappa(self):
        with tempfile.TemporaryDirectory() as td:
            src = Path(td) / "in.xlsx"
            dst = Path(td) / "out.xlsx"
            cego = Path(td) / "cego.xlsx"
            _wb_minimo(src)
            _wb_minimo(cego, nuclear=[True, False, True, False])
            plano = Path(td) / "plano.yaml"
            plano.write_text("desduplicacao: nenhuma\n", encoding="utf-8")
            ta.analisar(src, dst, desduplicacao="contexto",
                        plano_a_priori=plano, kappa_cego=cego)
            with pd.ExcelFile(dst) as xl:
                sheets = xl.sheet_names
            self.assertIn("16_Kappa", sheets)
            res = pd.read_excel(dst, sheet_name="1_Resumo")
            val = dict(zip(res["indicador"].astype(str), res["valor"]))
            self.assertEqual(str(val["Modo desduplicacao"]), "nenhuma")
            kap = pd.read_excel(dst, sheet_name="16_Kappa")
            self.assertIn("kappa", kap.columns)
            self.assertGreaterEqual(int(kap.loc[0, "n_emparelhados"]), 1)


if __name__ == "__main__":
    unittest.main()
