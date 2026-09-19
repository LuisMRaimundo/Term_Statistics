# -*- coding: utf-8 -*-
"""Gráfico nativo Excel de frequência (versão editável de _g_freq_token.png)."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

import textura_plots as tplot


class TestBarrasAgrupadasXlsx(unittest.TestCase):
    def test_escreve_dados_e_grafico_nativo(self):
        rotulos = ["homogen", "static", "uniform"]
        hits = [17, 12, 5]
        docs = [10, 8, 4]
        cores = [tplot.cor_canonical(t) for t in rotulos]
        with tempfile.TemporaryDirectory() as td:
            dest = Path(td) / "_g_freq_token.xlsx"
            out = tplot.barras_horizontais_agrupadas_xlsx(
                rotulos, hits, docs, dest,
                titulo="Frequência por termo canónico  (N_hits=34)",
                subtitulo="Hits nucleares em 8_Concordancia",
                xlabel="Ocorrências (N_hits)",
                rodape="TEXTURA  ·  pesquisa bibliográfica de termos",
                cores=cores,
            )
            self.assertEqual(out, dest)
            self.assertTrue(dest.is_file())

            wb = load_workbook(dest)
            ws = wb["Frequência"]
            self.assertIn("Frequência por termo canónico", str(ws["A1"].value))
            self.assertEqual(ws["A5"].value, "canonical_term")
            self.assertEqual(ws["B5"].value, "N_hits")
            self.assertEqual(ws["C5"].value, "n_documentos")
            self.assertEqual(ws["A6"].value, "homogen")
            self.assertEqual(int(ws["B6"].value), 17)
            self.assertEqual(int(ws["C6"].value), 10)
            self.assertEqual(ws["A8"].value, "uniform")

            self.assertEqual(len(ws._charts), 1)
            chart = ws._charts[0]
            self.assertEqual(chart.type, "bar")
            self.assertEqual(chart.grouping, "clustered")
            self.assertEqual(len(chart.series), 2)
            self.assertEqual(len(chart.series[0].dPt), 3)
            self.assertEqual(
                str(chart.series[0].dPt[0].spPr.solidFill.srgbClr),
                tplot._rgb(cores[0]),
            )
            self.assertEqual(
                str(chart.series[1].dPt[0].spPr.solidFill.srgbClr),
                tplot._rgb(tplot._clarear_hex(cores[0])),
            )
            self.assertEqual(chart.x_axis.scaling.orientation, "maxMin")
            wb.close()


if __name__ == "__main__":
    unittest.main()
